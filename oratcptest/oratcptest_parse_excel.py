import openpyxl
import pandas as pd
import plotly.express as px

# Defining the raw data
data = """
(14:36:44) The server is ready.
                    Throughput             Latency
(14:36:49)     17.515 Mbytes/s           57.093 ms
(14:36:54)     19.667 Mbytes/s           50.848 ms
(14:36:59)     17.948 Mbytes/s           55.717 ms
(14:37:04)     17.502 Mbytes/s           57.136 ms
(14:37:09)     18.577 Mbytes/s           53.831 ms
(14:37:14)     18.217 Mbytes/s           54.894 ms
(14:37:14) Test finished.
               Socket send buffer = 1048832 bytes
                  Avg. throughput = 18.234 Mbytes/s
                     Avg. latency = 54.843 ms
"""

# Parsing the data
lines = data.split("\n")
timestamps, throughputs, latencies = [], [], []
for line in lines:
    if 'Mbytes/s' in line and 'Throughput' not in line and 'Avg.' not in line:
        parts = line.split()
        timestamps.append(parts[0][1:-1])
        throughputs.append(float(parts[1]))
        latencies.append(float(parts[3].replace("ms", "")))

# Creating a DataFrame
df = pd.DataFrame({
    'Time': timestamps,
    'Throughput (Mbytes/s)': throughputs,
    'Latency (ms)': latencies
})

# Create a new Excel workbook and add the DataFrame to it
excel_path = 'Network_Performance_Metrics.xlsx'
with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
    df.to_excel(writer, sheet_name='Data', index=False)

# Open the existing Excel workbook and create a chart for Throughput and Latency
wb = openpyxl.load_workbook(excel_path)
ws = wb['Data']

# Chart for Throughput
chart_throughput = openpyxl.chart.LineChart()
chart_throughput.title = "Throughput over Time"
chart_throughput.style = 13
chart_throughput.x_axis.title = 'Time'
chart_throughput.y_axis.title = 'Throughput (Mbytes/s)'
data_throughput = openpyxl.chart.Reference(ws, min_col=2, min_row=1, max_col=2, max_row=len(throughputs) + 1)
cats = openpyxl.chart.Reference(ws, min_col=1, min_row=2, max_row=len(timestamps) + 1)
chart_throughput.add_data(data_throughput, titles_from_data=True)
chart_throughput.set_categories(cats)
ws.add_chart(chart_throughput, "E5")

# Chart for Latency
chart_latency = openpyxl.chart.LineChart()
chart_latency.title = "Latency over Time"
chart_latency.style = 13
chart_latency.x_axis.title = 'Time'
chart_latency.y_axis.title = 'Latency (ms)'
data_latency = openpyxl.chart.Reference(ws, min_col=3, min_row=1, max_col=3, max_row=len(latencies) + 1)
chart_latency.add_data(data_latency, titles_from_data=True)
chart_latency.set_categories(cats)
ws.add_chart(chart_latency, "M5")

# Save the workbook
wb.save(excel_path)
